from io import BytesIO

import pandas as pd
import pytest
from pypdf import PdfReader

from analyst_studio.services.analytics import auto_charts, custom_chart, insight_summary
from analyst_studio.services.export import csv_bytes, excel_bytes, pdf_bytes
from analyst_studio.services.insights import generate_insight


@pytest.fixture
def frame():
    return pd.DataFrame({"category": ["A", "A", "B"], "amount": [10, 20, 30],
                         "profit": [1, 2, 3], "date": ["2025-01-01", "2025-01-02", "2025-01-02"]})


def test_auto_charts_and_aggregate(frame):
    charts = auto_charts(frame)
    assert {c.kind for c in charts} == {"histogram", "bar", "line", "heatmap"}
    assert len(charts) == 5
    assert sum(next(c for c in charts if c.kind == "histogram").y) == 3
    assert next(c for c in charts if c.kind == "line").y == [1, 2]
    custom = custom_chart(frame, "Bar", "category", "amount")
    assert custom.y == [30, 30]
    assert "Pearson" in insight_summary(frame)


def test_empty_and_invalid_chart(frame):
    assert auto_charts(frame.iloc[:0]) == []
    with pytest.raises(ValueError):
        custom_chart(frame, "Scatter", "category", "amount")
    with pytest.raises(ValueError):
        custom_chart(frame, "Histogram", "category", "")


def test_line_chart_handles_mixed_timezone_offsets():
    frame = pd.DataFrame({"date": ["2025-01-01T12:00:00+07:00", "2025-01-02T12:00:00+00:00"], "amount": [1, 2]})
    chart = custom_chart(frame, "Line", "date", "amount")
    assert chart.y == [1, 2]


def test_data_exports(frame):
    assert pd.read_csv(BytesIO(csv_bytes(frame))).equals(frame)
    assert pd.read_excel(BytesIO(excel_bytes(frame))).equals(frame)


def test_excel_formula_is_text():
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(excel_bytes(pd.DataFrame({"text": ["=1+1"]}))))
    assert workbook.active["A2"].data_type == "s"


def test_excel_preserves_infinity_as_text_instead_of_blank():
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(excel_bytes(pd.DataFrame({"x": [float("inf"), float("-inf")]}))))
    assert workbook.active["A2"].value == "inf"
    assert workbook.active["A3"].value == "-inf"


def test_pdf_has_stats_charts_and_insight(frame):
    content = pdf_bytes(frame, auto_charts(frame), ["Trim name"], "AI summary example", "sample.csv", "All rows")
    assert content.startswith(b"%PDF")
    reader = PdfReader(BytesIO(content))
    text = "".join(page.extract_text() for page in reader.pages)
    assert "sample.csv" in text and "AI summary example" in text
    assert "Trim name" in text
    assert sum(len(page.images) for page in reader.pages) >= 1


@pytest.mark.asyncio
async def test_missing_ai_key_is_nonfatal(monkeypatch):
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    message, ok = await generate_insight("summary")
    assert not ok and "GEMINI_API_KEY" in message


@pytest.mark.asyncio
async def test_ai_failure_is_nonfatal(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-placeholder")
    def unavailable(*args, **kwargs):
        raise RuntimeError("service unavailable")
    monkeypatch.setattr("analyst_studio.services.insights.genai.Client", unavailable)
    message, ok = await generate_insight("summary")
    assert not ok and "ไม่สามารถ" in message
